"""Build the ten-minute group-meeting deck (P1, NICE) on the group's PowerPoint template.

Usage: python3 build_groupmeeting.py <template.pptx> <out.pptx>
The template is the group template as carried by the advisor deck (P1_NICE_导师汇报_20261008.pptx): its first slide
(title page) is kept and edited, its other slides are removed, and every new slide uses the template's single layout.
Nine content slides (problem; cut finite elements; NICE; network; properties and accuracy; cost and scale; design;
conclusions; progress), a closing slide and three backup slides. Helpers (text boxes, cards, stat callouts, tables,
pictures) come from build_talk.py; display formulas are typeset with LaTeX (Latin Modern, as in the paper) and inserted
as pictures; the two-dimensional cut finite element schematic is drawn here. Numbers follow the R7 manuscript
(docs/paper_p1/MANUSCRIPT_CN.md); figures come from docs/paper_p1/figures. The speaker notes come from
groupmeeting_notes.py: a timed script (about ten minutes in all) and backup items on every slide; the same notes are also
written next to the deck as <deck>_讲稿.md."""
import sys, hashlib, subprocess, tempfile
from pathlib import Path
from PIL import Image
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

sys.path.insert(0, str(Path(__file__).resolve().parent))
import re  # noqa: E402
import build_talk  # noqa: E402
from groupmeeting_notes import render as notes, export_markdown  # noqa: E402
from build_talk import (rgb, font, text, card, stat, image, table, cite, Deck,  # noqa: E402
                        RED, DARK, GRAY, BLUE, LIGHT, LBLUE, X0)


def para_runs(p, spec, size, color, bold):
    """As build_talk.para_runs, but a run with 'ss' also gets _{..} / ^{..} as sub- and superscripts in the body font."""
    if isinstance(spec, str):
        spec = [spec]
    for item in spec:
        txt, o = (item, {}) if isinstance(item, str) else item
        pieces = re.split(r'(_\{[^}]*\}|\^\{[^}]*\})', txt) if (o.get('math') or o.get('ss')) else [txt]
        for piece in pieces:
            if not piece:
                continue
            base = 0
            if piece.startswith('_{'):
                piece, base = piece[2:-1], -25000
            elif piece.startswith('^{'):
                piece, base = piece[2:-1], 30000
            r = p.add_run(); r.text = piece
            font(r, o.get('size', size), o.get('color', color), o.get('bold', bold), o.get('italic', False), o.get('math', False))
            if base:
                r._r.get_or_add_rPr().set('baseline', str(base))


build_talk.para_runs = para_runs


def ss(s, **kw):
    return (s, dict(ss=True, **kw))


# ---------------------------------------------------------------------------------------------------------------
# Visual system (flat): the template supplies the red title and the footer; content uses one ink colour, one navy
# for numbers and table heads, the template red only for small heads, bullets and the lead word of a takeaway, and
# flat light panels without shadows.
from pptx.oxml.ns import qn  # noqa: E402
from lxml import etree  # noqa: E402

NAVY, INK, MUTED = '1F3864', '262626', '6B7280'
PANEL, PANEL_B, PANEL_R, RULE = 'F4F6F9', 'EDF2F8', 'FBF1F1', 'D9DDE3'
BLUE = RED                    # lead words of takeaways and secondary heads use the template red
LIGHT, LBLUE = PANEL, PANEL_B
GRAY = MUTED


def flat(shape, fill):
    """Solid fill, no outline, no theme style (hence no shadow)."""
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill); shape.line.fill.background()
    return shape


def box(slide, x, y, w, h, fill=PANEL, radius=0.035):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.adjustments[0] = radius
    return flat(b, fill)


def my_bullet(p, level=0, char='•'):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(Inches(0.24 + 0.24 * level))))
    pPr.set('indent', str(int(-Inches(0.22))))
    for tag in ('a:buClr', 'a:buSzPct', 'a:buFont', 'a:buNone', 'a:buChar', 'a:buAutoNum'):
        el = pPr.find(qn(tag))
        if el is not None:
            pPr.remove(el)
    bc = etree.SubElement(pPr, qn('a:buClr')); etree.SubElement(bc, qn('a:srgbClr')).set('val', RED if level == 0 else MUTED)
    etree.SubElement(pPr, qn('a:buSzPct')).set('val', '70000')
    etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
    etree.SubElement(pPr, qn('a:buChar')).set('char', '■' if level == 0 else '–')


build_talk.bullet = my_bullet
_text = build_talk.text


def text(slide, x, y, w, h, paras, size=16, color=INK, line=1.12, fill=None, margin=0.06, **kw):
    if fill:
        box(slide, x, y, w, h, fill)
        fill = None
        margin = max(margin, 0.18)
    return _text(slide, x, y, w, h, paras, size=size, color=color, line=line, fill=fill, margin=margin, **kw)


def card(slide, x, y, w, h, head, body, fill=PANEL, head_color=RED, size=15, head_size=17, name=None):
    box(slide, x, y, w, h, PANEL if fill in (PANEL, 'F2F2F2') else (PANEL_B if fill in (PANEL_B, 'EAF2FB') else fill))
    text(slide, x + 0.2, y + 0.14, w - 0.4, 0.45, [head], size=head_size, color=head_color if head_color != '0070C0' else RED,
         bold=True)
    if body:
        text(slide, x + 0.2, y + 0.62, w - 0.4, h - 0.75, body, size=size, space=6)


def stat(slide, x, y, w, number, label, color=NAVY, nsize=34, lsize=13, fill=PANEL, h=1.45):
    if fill:
        box(slide, x, y, w, h, PANEL)
    text(slide, x + 0.1, y + 0.12, w - 0.2, 0.7, [number], size=nsize, color=NAVY if color == RED else color, bold=True,
         align='c', anchor='m')
    text(slide, x + 0.2, y + 0.84, w - 0.4, h - 0.9, [label], size=lsize, color=MUTED, align='c')


def banner(slide, y, runs, h=0.75, fill=None, size=16):
    box(slide, X0, y, W, h, PANEL_R)
    return _text(slide, X0 + 0.2, y, W - 0.4, h, [runs], size=size, color=INK, anchor='m', margin=0.02, line=1.12)


def cite(slide, s):
    _text(slide, X0, 6.8, W, 0.3, [s], size=10, color=MUTED, margin=0.0)


def _cell_border(cell, sides, color=RULE, w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for side in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        old = tcPr.find(qn(side))
        if old is not None:
            tcPr.remove(old)
    for side in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        ln = etree.SubElement(tcPr, qn(side))
        if side in sides:
            ln.set('w', str(w)); sf = etree.SubElement(ln, qn('a:solidFill'))
            etree.SubElement(sf, qn('a:srgbClr')).set('val', color)
        else:
            ln.set('w', '0'); etree.SubElement(ln, qn('a:noFill'))
    # schema order: borders come before the cell fill
    fill = tcPr.find(qn('a:solidFill'))
    if fill is not None:
        tcPr.remove(fill); tcPr.append(fill)


def table(slide, x, y, w, rows, widths, size=13, head_fill=NAVY, row_h=0.4, hl_rows=(), hl_fill=PANEL_R, name=None):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for k in ('bandRow', 'firstRow'):
        tblPr.set(k, '0')
    st = tblPr.find(qn('a:tableStyleId'))
    if st is not None:
        st.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'       # "No Style, No Grid"
    for j, fw in enumerate(widths):
        tbl.columns[j].width = Inches(w * fw / sum(widths))
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = build_talk.MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]; c.text_frame.word_wrap = True
            head = i == 0
            para_runs(p, val, size, 'FFFFFF' if head else INK, head)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(head_fill if head else (hl_fill if i in hl_rows else 'FFFFFF'))
            _cell_border(c, () if head else ('a:lnB',))
    return gt

FIG = Path(__file__).resolve().parents[2] / 'paper_p1' / 'figures'
W = 12.33                                   # content width (inches)
EQ_DIR = Path(tempfile.mkdtemp(prefix='gm_eq_'))
DPI = 600

PREAMBLE = r"""\documentclass[border=2pt]{standalone}
\usepackage{lmodern}
\usepackage{amsmath,amssymb,bm}
\usepackage[HTML]{xcolor}
\begin{document}
"""


def eq_png(tex, size):
    """Typeset one display formula at `size` pt; return (png path, width in, height in)."""
    key = hashlib.md5(f'{size}|{tex}'.encode()).hexdigest()[:12]
    png = EQ_DIR / f'{key}.png'
    if not png.exists():
        src = EQ_DIR / f'{key}.tex'
        src.write_text(PREAMBLE + r'\fontsize{%g}{%g}\selectfont\color[HTML]{262626}$\displaystyle %s$' % (size, 1.2 * size, tex)
                       + '\n\\end{document}\n', encoding='utf-8')
        r = subprocess.run(['lualatex', '-interaction=nonstopmode', '-halt-on-error', src.name], cwd=EQ_DIR,
                           capture_output=True, text=True)
        if r.returncode:
            raise RuntimeError('LaTeX failed for: ' + tex + '\n' + r.stdout[-1500:])
        subprocess.run(['pdftocairo', '-png', '-transp', '-r', str(DPI), '-singlefile', f'{key}.pdf', key], cwd=EQ_DIR,
                       check=True)
    w, h = Image.open(png).size
    return png, w / DPI, h / DPI


def formula(slide, tex, x, y, w, h, size=22, fill=LIGHT, align='c'):
    """A rounded box (x, y, w, h) holding one formula at its natural size, centred vertically."""
    if fill:
        box(slide, x, y, w, h, fill)
    sz = size
    png, pw, ph = eq_png(tex, sz)
    while (pw > w - 0.2 or ph > h - 0.08) and sz > 16:
        sz -= 1
        png, pw, ph = eq_png(tex, sz)
    if pw > w - 0.2 or ph > h - 0.08:
        raise ValueError(f'formula too large for its box ({pw:.2f} x {ph:.2f} in): {tex}')
    if sz != size:
        print(f'  formula shrunk {size} -> {sz} pt: {tex[:60]}')
    px = x + (w - pw) / 2 if align == 'c' else x + 0.2
    slide.shapes.add_picture(str(png), Inches(px), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def label(head, body, color=RED):
    return [(head, {'bold': True, 'color': color})] + (list(body) if isinstance(body, list) else [(body, {})])


def m(s):
    return (s, {'math': True})


# ---------------------------------------------------------------------------------------------------------------
# Two-dimensional CutFEM schematic (drawn here; the paper has no such figure). A Schwarz-P section
# |cos 2pi x + cos 2pi y + C0| <= tau(x, y) with bilinear corner thicknesses, one planar cut, an N x N background grid.
# Active elements, cut-plane elements, ghost faces and retained nodes follow the definitions of Section 2.2-2.3 and
# Appendix A.1, evaluated on a raster. The legend is drawn on the slide in the deck font.
import numpy as np  # noqa: E402

SC_C0 = 0.22
SC_TAU = {(0, 0): 0.29, (1, 0): 0.35, (0, 1): 0.33, (1, 1): 0.43}
SC_N = np.array([np.cos(np.radians(12)), np.sin(np.radians(12))])
SC_B = 0.80
SC_COL = dict(mat='#7FA5CD', removed='#D5DAE0', active='#E9EFF6', cutel='#F2E6CC', grid='#B9C1CA', ghost='#C8393B',
              cut='#8C6A2F', frame='#243447', face='#2E5C8A', inner='#9AA4AE')


def _sc_band(x, y):
    t = (SC_TAU[0, 0] * (1 - x) * (1 - y) + SC_TAU[1, 0] * x * (1 - y) + SC_TAU[0, 1] * (1 - x) * y
         + SC_TAU[1, 1] * x * y)
    return np.abs(np.cos(2 * np.pi * x) + np.cos(2 * np.pi * y) + SC_C0) <= t


def _sc_keep(x, y):
    return SC_N[0] * x + SC_N[1] * y <= SC_B


def cutfem_schematic(N=8, px=1500, ss=3):
    """Render the schematic to a PNG in EQ_DIR; return its path."""
    S = px * ss // N * N; M = int(0.035 * S); T = S + 2 * M; h = S // N
    xs = ((np.arange(T) + 0.5 - M) / S).astype(np.float32)
    ys = (1 - (np.arange(T) + 0.5 - M) / S).astype(np.float32)
    X, Y = xs[None, :], ys[:, None]
    insq = (X >= 0) & (X <= 1) & (Y >= 0) & (Y <= 1)
    B = _sc_band(X, Y) & insq
    K = _sc_keep(X, Y)
    mk, mc = B & K, B & ~K
    core = mk[M:M + S, M:M + S].reshape(N, h, N, h)
    act, full = core.any(axis=(1, 3)), core.all(axis=(1, 3))          # [row from the top, column]
    # cut-plane elements: the cut plane passes through material inside the element
    d = np.array([-SC_N[1], SC_N[0]])
    pts = SC_N * SC_B + (np.linspace(-1.5, 1.5, 30001))[:, None] * d[None, :]
    pts = pts[(pts[:, 0] >= 0) & (pts[:, 0] <= 1) & (pts[:, 1] >= 0) & (pts[:, 1] <= 1)]
    cutel = np.zeros((N, N), bool)
    for x, y in pts[_sc_band(pts[:, 0], pts[:, 1])]:
        cutel[N - 1 - min(int(y * N), N - 1), min(int(x * N), N - 1)] = True
    cutel &= act
    # Q2 nodes (2N+1 per axis): node (a, b) sits at (a/2N, b/2N)
    active_nodes, ret_face, ret_cut = set(), set(), set()
    s = np.linspace(0, 1, 401)
    for r in range(N):
        for i in range(N):
            if not act[r, i]:
                continue
            j = N - 1 - r
            nodes = {(2 * i + a, 2 * j + b) for a in range(3) for b in range(3)}
            active_nodes |= nodes
            if cutel[r, i]:
                ret_cut |= nodes
            for side, on in (('L', i == 0), ('R', i == N - 1), ('B', j == 0), ('T', j == N - 1)):
                if not on:
                    continue
                if side in 'LR':
                    ex, ey = np.full_like(s, 0.0 if side == 'L' else 1.0), (j + s) / N
                    fn = {(0 if side == 'L' else 2 * N, 2 * j + b) for b in range(3)}
                else:
                    ex, ey = (i + s) / N, np.full_like(s, 0.0 if side == 'B' else 1.0)
                    fn = {(2 * i + a, 0 if side == 'B' else 2 * N) for a in range(3)}
                if (_sc_band(ex, ey) & _sc_keep(ex, ey)).any():          # a cell-boundary face that carries material
                    ret_face |= fn
    interior = active_nodes - ret_face - ret_cut
    ghosts = []                                    # interior faces between two active elements, not both full
    for r in range(N):
        for i in range(N):
            j = N - 1 - r
            if i + 1 < N and act[r, i] and act[r, i + 1] and not (full[r, i] and full[r, i + 1]):
                ghosts.append(((i + 1) / N, j / N, (i + 1) / N, (j + 1) / N))
            if r + 1 < N and act[r, i] and act[r + 1, i] and not (full[r, i] and full[r + 1, i]):
                ghosts.append((i / N, j / N, (i + 1) / N, j / N))

    def P(x, y):
        return (M + x * S, M + (1 - y) * S)

    img = np.full((T, T, 3), 255, np.uint8)

    def paint(mask, col):
        img[mask] = np.array([int(col[k:k + 2], 16) for k in (1, 3, 5)], np.uint8)

    el, elc = np.zeros((T, T), bool), np.zeros((T, T), bool)
    for r in range(N):
        for i in range(N):
            if act[r, i]:
                (elc if cutel[r, i] else el)[M + r * h:M + (r + 1) * h, M + i * h:M + (i + 1) * h] = True
    paint(el, SC_COL['active']); paint(elc, SC_COL['cutel']); paint(mc, SC_COL['removed']); paint(mk, SC_COL['mat'])
    im = Image.fromarray(img)
    from PIL import ImageDraw
    dr = ImageDraw.Draw(im)
    lw = max(2, S // 450)
    for k in range(N + 1):
        dr.line([P(k / N, 0), P(k / N, 1)], fill=SC_COL['grid'], width=lw)
        dr.line([P(0, k / N), P(1, k / N)], fill=SC_COL['grid'], width=lw)
    for x0, y0, x1, y1 in ghosts:
        dr.line([P(x0, y0), P(x1, y1)], fill=SC_COL['ghost'], width=int(1.7 * lw))
    dr.line([P(*pts[0]), P(*pts[-1])], fill=SC_COL['cut'], width=int(4 * lw))
    dr.rectangle([P(0, 1), P(1, 0)], outline=SC_COL['frame'], width=int(3 * lw))

    def dot(a, b, rad, fill, edge=None):
        cx, cy = P(a / (2 * N), b / (2 * N))
        dr.ellipse([cx - rad, cy - rad, cx + rad, cy + rad], fill=fill, outline=edge, width=int(1.2 * lw) if edge else 0)

    for a, b in interior:
        dot(a, b, S * 0.0048, SC_COL['inner'])
    for a, b in ret_face:
        dot(a, b, S * 0.0105, SC_COL['face'], 'white')
    for a, b in ret_cut:
        dot(a, b, S * 0.0105, SC_COL['cut'], 'white')
    out = EQ_DIR / 'cutfem_schematic.png'
    im.resize((T // ss, T // ss), Image.LANCZOS).save(out)
    return out


def legend_item(slide, x, y, kind, col, label_text, size=11):
    """One legend entry: a swatch ('fill', 'cell', 'line', 'dot', 'small') and its label, 0.27 in tall."""
    if kind in ('fill', 'cell'):
        sw = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.06), Inches(0.22), Inches(0.16))
        flat(sw, col.lstrip('#'))
        if kind == 'cell':
            sw.line.color.rgb = rgb(SC_COL['grid'].lstrip('#')); sw.line.width = Pt(0.75)
    elif kind == 'line':
        flat(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.115), Inches(0.24), Inches(0.05)),
             col.lstrip('#'))
    else:
        d = 0.15 if kind == 'dot' else 0.08
        flat(slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.11 - d / 2), Inches(y + 0.14 - d / 2), Inches(d), Inches(d)),
             col.lstrip('#'))
    _text(slide, x + 0.3, y, 1.28, 0.28, [label_text], size=size, color=INK, anchor='m', margin=0.0)


def ladder(slide, x, y, w, h, head, sub, n1, l1, n2, l2, fill=PANEL):
    """A result box: head, context line, main number with label, rule, second number with label."""
    box(slide, x, y, w, h, fill)
    text(slide, x + 0.08, y + 0.1, w - 0.16, 0.36, [head], size=15, color=INK, bold=True, align='c')
    text(slide, x + 0.08, y + 0.46, w - 0.16, 0.32, [sub], size=12, color=MUTED, align='c')
    text(slide, x + 0.04, y + 0.82, w - 0.08, 0.62, [n1], size=30, color=NAVY, bold=True, align='c', anchor='m')
    text(slide, x + 0.08, y + 1.44, w - 0.16, 0.32, [l1], size=13, color=MUTED, align='c')
    flat(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.4), Inches(y + 1.86), Inches(w - 0.8), Inches(0.014)), RULE)
    text(slide, x + 0.04, y + 1.96, w - 0.08, 0.52, [n2], size=22, color=NAVY, bold=True, align='c', anchor='m')
    text(slide, x + 0.08, y + 2.48, w - 0.16, 0.5, [l2], size=13, color=MUTED, align='c')


def chevron(slide, x, y, w=0.2, h=0.34, col=None):
    flat(slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h)), col or RULE)


def build(template, out):
    D = Deck(template)

    # ------------------------------------------------------------------ 1 title page (template slide, edited)
    s1 = D.prs.slides[0]
    for sh in s1.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith('P1 进展汇报') or t.startswith('组会汇报'):
            r = sh.text_frame.paragraphs[0].runs
            r[0].text = '带平衡校正的神经初始化静力凝聚（NICE）'
            for extra in r[1:]:
                extra.text = ''
            for rr in r:
                rr.font.size = Pt(34)
        elif t.startswith('切割薄壁'):
            r = sh.text_frame.paragraphs[0].runs
            r[0].text = '及其在切割 TPMS 点阵厚度设计中的应用　·　组会汇报　·　2026 年 10 月'
            for extra in r[1:]:
                extra.text = ''
    s1.notes_slide.notes_text_frame.text = notes('title')

    # ------------------------------------------------------------------ 2 problem
    s = D.slide('切割点阵的厚度设计：均匀化不准，细尺度分析太慢', notes('problem'))
    image(s, FIG / 'F00_problem.png', X0, 1.08, W, 3.45)
    Wc = (W - 2 * 0.3) / 3
    for k, (num, lab) in enumerate([
            ('27%', '均匀化：切割面固支板上低估初始柔度'),
            ('14–20 min', '细尺度精确分析：八胞元点阵分析一次（16 线程），每次迭代都要重做'),
            ('2,679–45,900', '各胞元保留的主自由度：切割面上的支撑和载荷作用于切割平面单元的全部节点')]):
        stat(s, X0 + k * (Wc + 0.3), 4.64, Wc, num, lab, nsize=30, lsize=13, h=1.42)
    banner(s, 6.16, [('思路：', {'bold': True, 'color': RED}),
                     ('保留静力凝聚的框架和全部主自由度，只近似内部位移，免去内部刚度矩阵的分解', {})], h=0.56)

    # ------------------------------------------------------------------ 3 CutFEM and discretisation
    s = D.slide('切割有限元离散：几何改变无需重新划分网格', notes('cutfem'))
    image(s, cutfem_schematic(), X0, 1.08, 4.72, 4.72)
    legend = [('fill', SC_COL['mat'], '材料层'), ('fill', SC_COL['removed'], '切除部分'), ('line', SC_COL['cut'], '切割平面'),
              ('cell', SC_COL['active'], '激活单元'), ('fill', SC_COL['cutel'], '切割平面单元'), ('line', SC_COL['ghost'], 'ghost 面'),
              ('dot', SC_COL['face'], '边界面主节点'), ('dot', SC_COL['cut'], '切割单元主节点'), ('small', SC_COL['inner'], '内部节点')]
    for k, (kind, col, lab) in enumerate(legend):
        legend_item(s, X0 + 0.05 + (k % 3) * 1.6, 5.88 + (k // 3) * 0.27, kind, col, lab)
    image(s, FIG / 'F08_geometry.png', 5.55, 1.08, 7.28, 2.48)
    text(s, 5.55, 3.72, 7.28, 3.0, [
        label('几何　', 'Schwarz-P 材料层；8 个角点厚度参数插值；至多一个平面切割'),
        label('网格　', '固定笛卡尔背景网格，每轴 32 个 Q2 单元（左图为每轴 8 个单元的二维示意）'),
        label('积分　', '在单元的材料部分上积分：子胞元细分，材料边界处裁剪四面体'),
        label('稳定　', [('ghost 罚项控制只含少量材料的单元（', {}), ss('γ = 10^{−4}'), ('）', {})]),
        label('主自由度　', '胞元边界面节点和切割平面单元的全部节点：2,679–45,900 个；内部自由度约为其 6 倍'),
    ], size=16, space=10)

    # ------------------------------------------------------------------ 4 NICE overview
    s = D.slide('NICE：保留全部主自由度，只近似内部位移', notes('nice'))
    image(s, FIG / 'F01_method_overview.png', X0, 1.08, 6.9, 5.62)
    formula(s, r'\widehat S=\widehat\Psi^{T}K\widehat\Psi,\qquad \widehat\Psi=\mathcal W\,\widehat\Psi_{\rm net}', 7.75, 1.12, 5.08,
            0.95, size=22, fill=PANEL_B)
    steps = [('网络', '由主自由度位移预测内部位移；以几何为输入，对 q 线性'),
             ('修正', '固定的两重网格修正 8 / Q1(17) / 8，减小内部不平衡力；不含训练参数'),
             ('刚度', '由修正后位移场的应变能计算；装配求解只需矩阵–向量乘积')]
    for k, (h_, b_) in enumerate(steps):
        y = 2.3 + k * 1.08
        _text(s, 7.75, y - 0.04, 0.55, 0.6, [str(k + 1)], size=28, color=RED, bold=True, margin=0.0)
        text(s, 8.3, y, 4.53, 0.38, [h_], size=17, color=INK, bold=True)
        text(s, 8.3, y + 0.38, 4.53, 0.66, [b_], size=15, color=INK)
    text(s, 7.75, 5.62, 5.08, 1.05, [[('不分解内部刚度矩阵，不形成稠密凝聚刚度矩阵；', {}),
                                      ('设计迭代中网络参数不变', {'bold': True, 'color': RED})]],
         size=15, fill=PANEL_R, anchor='m')

    # ------------------------------------------------------------------ 5 network architecture (the whole of Figure 3)
    # Figure 3 is cut into its two halves, (a, b) on the left and (c, d) on the right, shown at the same scale; the pixel
    # rows below bound the halves of F11_network_architecture.png (the bottom line on training and deployment is left out).
    s = D.slide('网络：几何分支算系数，位移分支对 q 线性', notes('network'))
    F3 = FIG / 'F11_network_architecture.png'
    fw, fh = Image.open(F3).size
    Wf = (W - 0.25) / 2
    for k, (top, bottom) in enumerate([(40, 1365), (1385, 2385)]):
        ar = (fw - 45) / (bottom - top)
        image(s, F3, X0 + k * (Wf + 0.25), 1.05, Wf, Wf / ar, crop=(45 / fw, top / fh, 0.0, 1 - bottom / fh))
    y_ab = 1.05 + Wf * (1365 - 40) / (fw - 45)                    # bottom of the (a, b) half
    y_cd = 1.05 + Wf * (2385 - 1385) / (fw - 45)                  # bottom of the (c, d) half
    text(s, X0, y_ab + 0.12, Wf, 6.72 - y_ab - 0.12, [
        [('单元特征 126 = 1 + 1 + 124　', {'bold': True, 'color': RED}),
         ('材料体积分数、其对数、124 个单元矩 ÷ 材料体积（第 3 页的 125 个矩去掉零阶）', {})],
        [('节点特征 11 = 4 + 1 + 6　', {'bold': True, 'color': RED}),
         ('4 个标记（主节点 / 胞元边界面 / 切割平面单元 / 弱连接）、1 个刚度量（对角刚度块范数的对数）、6 个位置项 sin、cos(2πx, 2πy, 2πz)', {})],
    ], size=14, space=8, fill=PANEL_B, anchor='m')
    text(s, X0 + Wf + 0.25, y_cd + 0.1, Wf, 6.72 - y_cd - 0.1, [
        label('几何分支 (a)　', '非线性，每个几何算一次，算出各运算的系数'),
        label('位移分支 (b)　', '对 q 线性；刚体运动绕过网络，主自由度在输出端重新赋值'),
        label('(c) 传得远　', '每轴 65/33/17/9 的 U 形多层网格，负责远距离耦合'),
        label('(d) 看得细　', '27 节点模板上的局部交互，负责单元尺度的几何细节'),
        label('规模无关　', [('权重不对应具体自由度：约 ', {}), ss('6×10^{5}'),
                         (' 参数的同一网络适用于 2,679–45,900 个主自由度', {})]),
    ], size=14, space=7)

    # ------------------------------------------------------------------ 6 properties and accuracy
    s = D.slide('性质由构造保证；从单胞到点阵，误差不累积', notes('properties'))
    formula(s, r'\boxed{\;\widehat S-S=H^{T}\!AH\succeq 0\;}', X0, 1.1, 4.55, 1.25, size=28, fill=PANEL_B)
    text(s, X0, 2.42, 4.55, 0.95, [[('命题 1（Ritz 恒等式）', {'bold': True, 'color': INK}),
                                    ('：位移恢复在主自由度上等于给定位移、精确再现刚体运动时成立，与网络参数无关', {})]],
         size=13, color=MUTED)
    Wp = (12.83 - 5.3 - 2 * 0.2) / 3
    for k, (h_, b_) in enumerate([
            ('可组装', ['对称半正定，只有刚体零能模态；不低于精确值，误差为二阶小量']),
            ('可改进', ['修正不增大误差；部署时可增加修正步数，无需重新训练']),
            ('可追溯', ['柔度误差不超过 β（能量误差按应变能占比加权）；灵敏度误差含一阶项，两者分别检验'])]):
        card(s, 5.3 + k * (Wp + 0.2), 1.1, Wp, 2.27, h_, b_, size=14, head_size=17)
    Wl = (W - 3 * 0.3) / 4
    boxes = [('单个胞元', '80 个验证几何，面力载荷', '0.074%', '平均能量误差', '6.89% → 0.0965%', '基础网络只在部署时加修正'),
             ('两胞元装配', '32 种构型，一个学习胞元', '≤0.27%', '柔度误差', '≤1.49%', '灵敏度误差'),
             ('两个八胞元点阵', '全部为学习胞元，均未参与训练', '<0.015%', '柔度误差', '≤0.09%', '点阵梯度误差'),
             ('柔度准 ≠ 灵敏度准', '例：U1/x，载荷加在相邻胞元上', '0.000073%', '柔度误差', '0.598%',
              '灵敏度误差，约为柔度误差的 8,200 倍')]
    for k, bx in enumerate(boxes):
        x = X0 + k * (Wl + 0.3)
        ladder(s, x, 3.6, Wl, 3.05, *bx, fill=PANEL_R if k == 3 else PANEL)
        if k < 2:
            chevron(s, x + Wl + 0.05, 4.82, col=MUTED)

    # ------------------------------------------------------------------ 7 cost and scale
    s = D.slide('成本：快约十倍，耗时随胞元数线性增长', notes('cost'))
    cd = CategoryChartData()
    cd.categories = ['八胞元块体', '八胞元单层']
    cd.add_series('整体点阵直接求解（16 线程）', (868, 1042))
    cd.add_series('常规精确静力凝聚（16 线程）', (950, 1218))
    cd.add_series('NICE（一块 GPU，含灵敏度）', (79.2, 105.5))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(X0), Inches(1.08), Inches(6.2), Inches(3.25), cd)
    ch = gf.chart
    ch.has_title = True; ch.chart_title.text_frame.text = '一次分析的耗时（s）'
    for r in ch.chart_title.text_frame.paragraphs[0].runs:
        font(r, 14, INK, True)
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(11)
    for ser, col in zip(ch.plots[0].series, ('8A94A0', 'C9CFD6', RED)):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
    pl = ch.plots[0]; pl.has_data_labels = True; pl.data_labels.font.size = Pt(11)
    pl.data_labels.number_format = '#,##0'; pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.gap_width = 70
    va = ch.value_axis; va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = rgb('E5E7EB')
    va.tick_labels.font.size = Pt(10); va.format.line.fill.background(); va.maximum_scale = 1400
    ch.category_axis.tick_labels.font.size = Pt(12)
    image(s, FIG / 'F16_scale_plates.png', 6.95, 1.08, 5.88, 3.25)
    Ws = (W - 3 * 0.25) / 4
    for k, (num, lab) in enumerate([('79–105 s', 'NICE 一次分析（含灵敏度）'),
                                    ('9.2 GiB', 'GPU 内存峰值（直接求解 71–86 GiB）'),
                                    ('10.1–10.9 s', '每胞元每次设计迭代（24–110 胞元）'),
                                    ('19.4 min', '110 胞元板一次迭代（3270 万自由度）')]):
        stat(s, X0 + k * (Ws + 0.25), 4.48, Ws, num, lab, nsize=28, lsize=13, h=1.45)
    banner(s, 6.05, [('快的原因：', {'bold': True, 'color': RED}),
                     ('内部位移由一次网络计算和固定步数的修正得到，不分解内部刚度矩阵；装配问题只在主自由度上求解', {})], h=0.6)
    cite(s, '一块 GPU 对 16 个 CPU 线程，比较的是各自部署方式下的成本；均不含数据生成和训练')

    # ------------------------------------------------------------------ 8 thickness design
    s = D.slide('厚度设计：复现精确静力凝聚的结果', notes('design'))
    image(s, FIG / 'F13_optimisation.png', X0, 1.08, 7.75, 3.72)
    image(s, FIG / 'F14_plate_design3d.png', 8.5, 1.08, 4.33, 3.2, crop=(0.0, 0.0, 0.5, 0.47))
    text(s, 8.5, 4.3, 4.33, 0.4, ['切割面固支板的 NICE 最终设计'], size=12, color=MUTED, align='c')
    for k, (num, lab) in enumerate([('≤0.0051', '算例 A：与精确静力凝聚设计的角点参数之差'),
                                    ('27%', '固支板：均匀化低估初始设计柔度（加密后 25.8%）'),
                                    ('0.03% / 0.3%', 'NICE 最终设计的柔度 / 梯度偏差'),
                                    ('+2.44%', '均匀化设计的精确柔度（材料少 0.18%）')]):
        stat(s, X0 + k * (Ws + 0.25), 5.0, Ws, num, lab, nsize=28, lsize=13, h=1.5)

    # ------------------------------------------------------------------ 9 conclusions and limitations
    s = D.slide('结论与局限', notes('conclusions'))
    for k, (h_, b_) in enumerate([
            ('只近似内部位移', '保留全部主自由度：切割面上的支撑和载荷直接施加；同一网络适用于随切割变化的主自由度集合'),
            ('性质由构造保证', '对称半正定、不低于精确值、误差为二阶小量，与网络精度无关；修正不增大误差'),
            ('柔度与灵敏度分别检验', '八胞元点阵柔度误差 <0.015%、梯度 ≤0.09%；分析快约十倍；设计复现精确静力凝聚的结果')]):
        x = X0 + k * (Wc + 0.3)
        box(s, x, 1.15, Wc, 3.0, PANEL)
        _text(s, x + 0.22, 1.25, 1.0, 0.6, ['0%d' % (k + 1)], size=30, color=RED, bold=True, margin=0.0)
        text(s, x + 0.2, 1.9, Wc - 0.4, 0.45, [h_], size=18, color=INK, bold=True)
        text(s, x + 0.2, 2.42, Wc - 0.4, 1.65, [b_], size=15, line=1.18)
    text(s, X0, 4.35, W, 0.45, ['局限'], size=19, color=INK, bold=True)
    for k, (h_, b_) in enumerate([
            ('适用范围', '一种胞族（Schwarz-P）、单平面切割、一种分辨率；新胞族需重新训练'),
            ('误差', '学习模型偏刚，柔度被低估，不是保守估计；没有经认证的误差界'),
            ('规模', [('受预条件子中 ', {}), ss('K_{PP}'), (' 直接分解的限制', {})])]):
        card(s, X0 + k * (Wc + 0.3), 4.88, Wc, 1.78, h_, [b_], fill=PANEL_R, size=15, head_size=17)

    # ------------------------------------------------------------------ 10 progress and plan
    s = D.slide('进展与计划', notes('plan'))
    Wf = (W - 2 * 0.45) / 3
    for k, (h_, b_, f_) in enumerate([('已完成', '中英文正文与补充材料修订完成；投稿信、亮点已起草', PANEL_B),
                                       ('待办', '作者信息与各项声明；图形摘要更新；通读定稿', PANEL),
                                       ('投稿', 'CMAME，同时挂 arXiv', PANEL_R)]):
        x = X0 + k * (Wf + 0.45)
        card(s, x, 1.2, Wf, 1.7, h_, [b_], fill=f_, size=17, head_size=19, head_color=NAVY if k == 0 else RED)
        if k < 2:
            chevron(s, x + Wf + 0.12, 1.88, col=MUTED)
    text(s, X0, 3.2, W, 0.45, ['后续工作'], size=19, color=INK, bold=True)
    for k, (h_, b_) in enumerate([('P2　局部量驱动的设计', '以应力等局部量为目标或约束；配后验误差估计与自适应修正'),
                                   ('P3　跨胞族与几何泛化', '多胞族联合训练，在留出的胞族上检验；映射或曲面边界的胞元'),
                                   ('规模', '预条件子增加一层或采用迭代粗求解，扩大单 GPU 上可分析的规模')]):
        card(s, X0 + k * (Wc + 0.3), 3.75, Wc, 2.35, h_, [b_], size=17, head_size=18)

    # ------------------------------------------------------------------ 11 closing
    s = D.slide('谢谢！敬请各位老师和同学批评指正', notes('thanks'))
    image(s, FIG / 'F14_plate_design3d.png', 2.0, 1.3, 9.33, 4.7, crop=(0.0, 0.0, 0.5, 0.47))
    text(s, X0, 6.1, W, 0.5, ['切割面固支板的 NICE 最终设计（NICE 柔度 78.26，精确柔度 78.29）'], size=15, color=GRAY, align='c')

    # ------------------------------------------------------------------ backup 1: training and deployment
    s = D.slide('备用：训练与部署', notes('b_training'))
    formula(s, r'\mathcal L(\theta)=\frac1B\sum_{j=1}^{B}\log\!\big(q_j^{T}\widehat S q_j\big)'
               r'+w_s\frac{1}{|\mathcal J_s|}\sum_{j\in\mathcal J_s}\frac{\|\tilde{\bm s}_j-\bm s_j\|_2^2}{\|\bm s_j\|_2^2},'
               r'\qquad q_j^{T}Sq_j=1', X0, 1.2, W, 1.25, size=22)
    text(s, X0, 2.6, 5.9, 4.05, [
        ('-', '第一项：能量比的对数平均；由命题 1 非负，仅当恢复位移场精确时为零'),
        ('-', [('第二项：八分量厚度灵敏度相对参考值的误差（', {}), ss('w_{s} = 1'), ('）', {})]),
        ('-', '测试位移：多项式与多尺度位移，集中节点载荷、面力载荷、弹簧支撑下的响应，相邻胞元施加的边界位移'),
        ('-', '损失由修正后的算子计算，梯度经光滑和粗网格校正传到网络参数'),
    ], size=16, space=12)
    table(s, 6.65, 2.65, 6.18, [
        ['变体', '网络参数', '训练修正', '评估修正'],
        ['基础网络', '训练 40,000 步（305 个几何）', '无', '无'],
        ['基础网络加修正', '沿用基础网络', '无', '8/Q1(17)/8'],
        [[('NICE', {'bold': True, 'color': RED})], '继续训练 15,000 步（591 个几何的集合）', '8/Q1(17)/8', '8/Q1(17)/8'],
    ], (1.7, 2.5, 1.05, 1.05), size=13, row_h=0.62, hl_rows=(3,))
    text(s, 6.65, 5.3, 6.18, 1.35, [[('部署：', {'bold': True, 'color': RED}),
                                     ('每个几何只算一次几何编码和修正设置；每次乘以 Ŝ 执行网络、修正、K 及其转置；预条件子在每次设计迭代中分解一次 ', {}),
                                     m('K_{PP}')]], size=15, fill=LIGHT, margin=0.15, anchor='m')
    cite(s, '论文第 3.5、3.6 节，式 (13)，表 2')

    # ------------------------------------------------------------------ backup 2: correction effect
    s = D.slide('备用：修正的作用', notes('b_correction'))
    image(s, FIG / 'F04_correction.png', X0, 1.1, 7.3, 2.75, crop=(0.0, 0.0, 0.0, 0.5))
    table(s, X0, 3.95, 7.3, [
        ['胞元', '基础网络（无修正）', '零内部位移', '图调和', '基础网络', 'NICE'],
        ['M1', '13.5', '1097', '17.2', '0.186', '0.131'],
        ['M2', '3.64', '324', '6.72', '0.0273', '0.0358'],
        ['H1', '1.81', '39.9', '1.28', '0.0131', '0.0115'],
        ['H2', '35.0', '28.3', '0.0806', '0.0117', '0.0148'],
        ['U2', '0.402', '47.2', '1.23', '0.00424', '0.00465'],
    ], (0.6, 1.8, 1.25, 0.95, 1.2, 0.95), size=13, row_h=0.39)
    text(s, X0, 6.36, 7.3, 0.3, ['后四列为施加同一 8/Q1(17)/8 修正后的平均能量误差（%），面力载荷'], size=12, color=GRAY)
    text(s, 8.1, 1.2, 4.73, 5.45, [
        ('-', 'M1：13.5% → 八步光滑 4.68% → 粗网格校正后八步光滑 0.230% → 8 / Q1(17) / 8 0.186%'),
        ('-', '光滑消除高频分量，低频分量由粗网格校正处理'),
        ('-', '相同修正下，图调和初始场遗留的误差为 NICE 的 5.4–265 倍'),
        ('-', '基础网络只在部署时加修正：80 个验证几何的平均能量误差 6.89% → 0.0965%'),
    ], size=15, space=10)
    cite(s, '论文第 5.3、5.5 节，图 8(a,b)，表 4')

    # ------------------------------------------------------------------ backup 3: table 3
    s = D.slide('备用：全部装配算例与精确静力凝聚的偏差', notes('b_table3'))
    table(s, X0, 1.15, W, [
        ['算例', '每构型胞元数（切割）', '载荷或设计', '柔度误差 (%)', '灵敏度或梯度误差 (%)'],
        ['两胞元装配，七个选择胞元', '2（0 或 1）', '14 个构型，各六个面载荷', '≤ 0.056', '灵敏度 ≤ 0.67'],
        ['两胞元装配，切割面面力载荷', '2（1）', '10 个构型，各三个切割面面力载荷', '≤ 0.10', '灵敏度 ≤ 0.16'],
        ['两胞元装配，九个检查点选择之外的胞元', '2（0 或 1）', '18 个构型，各六个面载荷', '≤ 0.27', '灵敏度 ≤ 1.49'],
        ['全学习胞元点阵，面载荷', '8（4），8（3）', '三个面载荷', '≤ 0.015', '灵敏度 ≤ 0.144；梯度 0.04–0.09'],
        ['全学习胞元点阵，随机载荷', '8（4），8（3）', '三个随机载荷', '≤ 0.069', '灵敏度 ≤ 0.454；梯度 0.22–0.28'],
        ['算例 A，厚度优化', '8（4）', '第 0、12、23 次迭代的设计', '0.011–0.028', '梯度 0.069–0.33'],
        ['切割面固支板', '24（8）', '初始、最终、均匀化及两个均匀设计', '0.011–0.032', '梯度 0.029–0.32'],
        ['规模算例的板', '24（8），51（12）', '初始设计', '0.013，0.011', '梯度 0.045，0.047'],
    ], (3.5, 2.15, 3.0, 1.4, 2.55), size=13, row_h=0.5, hl_rows=(4, 6, 7))
    banner(s, 5.85, [('柔度误差：', {'bold': True, 'color': RED}), ('全部 ≤ 0.27%；', {}),
                     ('全学习胞元的点阵与设计：', {'bold': True, 'color': RED}), ('柔度 ≤ 0.069%，梯度 ≤ 0.33%', {})], h=0.7)
    cite(s, '论文表 3；两胞元装配中只有被测胞元为学习胞元，其余算例中所有胞元均为学习胞元')

    n = len(D.prs.slides)
    D.total(n)
    D.prs.save(out)
    print(out, n, 'slides')
    keys = ['title', 'problem', 'cutfem', 'nice', 'network', 'properties', 'cost', 'design', 'conclusions', 'plan', 'thanks',
            'b_training', 'b_correction', 'b_table3']
    assert len(keys) == n
    md = Path(out).with_name(Path(out).stem + '_讲稿.md')
    export_markdown(md, {k: i + 1 for i, k in enumerate(keys)})
    print(md)


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
